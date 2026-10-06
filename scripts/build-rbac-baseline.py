"""Turn the RBAC spreadsheet into the baseline JSON the API seeds using openpyxl."""

import json
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "docs" / "TK I PMS I New Modules RBAC I V01.xlsx"
OUT = ROOT / "apps" / "backend" / "Modules" / "Rbac" / "rbac-baseline.json"

ROLE_NAMES = {
    "CEO": "CEO",
    "COO": "COO",
    "CTO": "CTO",
    "IT Admin": "IT Admin",
    "Accounts": "Accounts",
    "HR": "HR",
    "Sales Manger": "Sales Manager",
    "Sales team member": "Sales team member",
    "PMO (Project Management Office)": "PMO",
    "EM (Engagement Manager)": "EngagementManager",
    "R&D -Team member": "R&D - Team member",
    "SOC - Team Member": "SOC-Team Member",
    "SOC - Team Leader": "SOC-Team Leader",
    "SOC - Manager": "SOC-Manager",
    "SOC - Senior Manager": "SOC-Senior Manager",
    "SOC - HOD": "SOC-HOD",
    "Consulting -Team member": "Consulting-Team member",
    "Consulting - Team Leader": "Consulting-Team Leader",
    "Consulting - Manager": "Consulting-Manager",
    "Consulting - Senior Manager": "Consulting-Senior Manager",
    "Consulting - HOD": "Consulting-HOD",
    "Testing - Team Member": "Testing-Team Member",
    "Testing - Team Leader": "Testing-Team Leader",
    "Testing - Manager": "Testing-Manager",
    "Testing - Senior Manager": "Testing Senior Manager",
    "Testing - HOD": "Testing HOD",
    "Intern": "Intern",
}


def parse_access(val: str) -> list[int]:
    raw = (val or "").strip().lower()
    if raw in ("read & write", "write", "rwe", "all rwe", "manage", "full manage", "r/w"):
        return [1, 1]
    if raw in ("read only", "read", "view only", "view", "ro"):
        return [1, 0]
    return [0, 0]  # "no access", empty, etc.


def main() -> None:
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    sheet = wb.active

    # Resolve merged cells
    merged_map = {}
    for rng in sheet.merged_cells.ranges:
        top_left_val = sheet.cell(row=rng.min_row, column=rng.min_col).value
        for row in range(rng.min_row, rng.max_row + 1):
            for col in range(rng.min_col, rng.max_col + 1):
                merged_map[(row, col)] = top_left_val

    def get_val(r: int, c: int):
        if (r, c) in merged_map:
            return merged_map[(r, c)]
        return sheet.cell(row=r, column=c).value

    # Extract roles (Rows 6 to 32)
    roles = []
    for r in range(6, 33):
        label = get_val(r, 1)
        if not label:
            continue
        clean_label = str(label).strip()
        if clean_label in ROLE_NAMES:
            roles.append((r, ROLE_NAMES[clean_label]))
        else:
            print(f"Warning: Unknown role label at row {r}: '{clean_label}'")

    # Extract leaves (Columns 3 to 51)
    leaves = []
    for c in range(3, 52):
        path = [get_val(r, c) for r in range(1, 6)]
        path_str = [str(p).strip() if p is not None else "" for p in path]
        if not any(path_str):
            continue

        grants = {}
        for row, role_name in roles:
            cell_val = get_val(row, c)
            raw = str(cell_val).strip() if cell_val is not None else "No Access"
            grants[role_name] = parse_access(raw)

        leaves.append(
            {
                "module": path_str[0] or None,
                "submodule": path_str[1] or None,
                "subSubmodule": path_str[2] or None,
                "widget": path_str[3] or None,
                "tab": path_str[4] or None,
                "grants": grants,
            }
        )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(
        json.dumps({"version": "excel-v02", "leaves": leaves}, indent=2),
        encoding="utf-8",
    )
    print(f"Successfully generated {len(leaves)} leaves for {len(roles)} roles to {OUT}")


if __name__ == "__main__":
    main()
